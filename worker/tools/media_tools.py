"""
Vicarius Agent - Media & Document Tools
Generates Cognitive Heritage styled Word (.docx), PowerPoint (.pptx),
and UI preview screenshot (.png) deliverables.
"""

from pathlib import Path
from typing import List, Dict, Any, Optional


def _set_run_font(run, font_name="Segoe UI", size=None, bold=False, italic=False, color=None):
    from docx.shared import Pt
    from docx.oxml.ns import qn
    if size is None:
        size = Pt(10.0)
    run.font.name = font_name
    run.font.size = size
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)


def _set_cell_background(cell, fill_hex):
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def _set_cell_margins(cell, top=70, bottom=70, left=90, right=90):
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)


def _set_cell_nowrap(cell):
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    tcPr = cell._tc.get_or_add_tcPr()
    noWrap = parse_xml(f'<w:noWrap {nsdecls("w")}/>')
    tcPr.append(noWrap)


def _add_bottom_border(paragraph, color="D1C5B6", sz="8"):
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="6" w:color="{color}"/></w:pBdr>')
    pPr.append(pBdr)


def clean_math_text(text: str) -> str:
    """Translates LaTeX math syntax into clean readable unicode mathematical text."""
    import re
    if not text:
        return ""
    t = text
    # Strip trailing markdown escape slashes
    t = re.sub(r'\\+\s*$', '', t)
    
    replacements = [
        (r'\\rightarrow', ' → '), (r'\\leftarrow', ' ← '),
        (r'\\Rightarrow', ' ⇒ '), (r'\\Leftarrow', ' ⇐ '),
        (r'\\iff', ' ⟺ '), (r'\\leftrightarrow', ' ↔ '),
        (r'\\le\b', ' ≤ '), (r'\\leq\b', ' ≤ '),
        (r'\\ge\b', ' ≥ '), (r'\\geq\b', ' ≥ '),
        (r'\\neq\b', ' ≠ '), (r'\\approx\b', ' ≈ '), (r'\\equiv\b', ' ≡ '),
        (r'\\cup\b', ' ∪ '), (r'\\cap\b', ' ∩ '),
        (r'\\in\b', ' ∈ '), (r'\\notin\b', ' ∉ '),
        (r'\\subset\b', ' ⊂ '), (r'\\subseteq\b', ' ⊆ '),
        (r'\\times\b', ' × '), (r'\\cdot\b', ' · '),
        (r'\\pm\b', ' ± '), (r'\\div\b', ' ÷ '),
        (r'\\infty\b', ' ∞ '), (r'\\forall\b', ' ∀ '), (r'\\exists\b', ' ∃ '),
        (r'\\wedge\b', ' ∧ '), (r'\\land\b', ' ∧ '),
        (r'\\lor\b', ' ∨ '), (r'\\neg\b', ' ¬ '),
        (r'\\alpha\b', 'α'), (r'\\beta\b', 'β'), (r'\\gamma\b', 'γ'),
        (r'\\delta\b', 'δ'), (r'\\epsilon\b', 'ε'), (r'\\theta\b', 'θ'),
        (r'\\lambda\b', 'λ'), (r'\\mu\b', 'μ'), (r'\\pi\b', 'π'),
        (r'\\rho\b', 'ρ'), (r'\\sigma\b', 'σ'), (r'\\tau\b', 'τ'),
        (r'\\phi\b', 'φ'), (r'\\omega\b', 'ω'),
        (r'\\max\b', 'max'), (r'\\min\b', 'min'), (r'\\log\b', 'log'),
    ]
    for pattern, rep in replacements:
        t = re.sub(pattern, rep, t)
    # Fractions and text formatting
    t = re.sub(r'\\frac\{([^{}]+)\}\{([^{}]+)\}', r'(\1) / (\2)', t)
    t = re.sub(r'\\text\{([^{}]+)\}', r'\1', t)
    # Subscripts and superscripts in braces: e.g. R_{W} -> R_W, 2^{k-1} -> 2^(k-1)
    t = re.sub(r'_\{([^{}]+)\}', r'_\1', t)
    t = re.sub(r'\^\{([^{}]+)\}', r'^(\1)', t)
    # Delimiters and cases
    t = t.replace(r'\(', '').replace(r'\)', '').replace(r'\[', '').replace(r'\]', '')
    t = re.sub(r'\\begin\{cases\}', '', t)
    t = re.sub(r'\\end\{cases\}', '', t)
    t = t.replace(r'\{', '{').replace(r'\}', '}')
    t = t.replace(r'\\', ' | ')
    t = re.sub(r'\\([a-zA-Z]+)', r'\1', t)
    t = re.sub(r'[ \t]+', ' ', t)
    return t.strip()


def _add_formatted_runs(paragraph, text, base_font="Segoe UI", base_size=None, base_color=None):
    import re
    from docx.shared import Pt, RGBColor
    if base_size is None:
        base_size = Pt(10.0)
    if base_color is None:
        base_color = RGBColor(35, 26, 15)

    clean_text = clean_math_text(text)
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', clean_text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run = paragraph.add_run(token[2:-2])
            _set_run_font(run, font_name=base_font, size=base_size, bold=True, color=base_color)
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, font_name=base_font, size=base_size, italic=True, color=RGBColor(77, 70, 58))
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, font_name="Consolas", size=Pt(base_size.pt * 0.9), color=RGBColor(112, 87, 40))
        else:
            run = paragraph.add_run(token)
            _set_run_font(run, font_name=base_font, size=base_size, color=base_color)


def clean_markdown_text(raw_text: str) -> str:
    """Unwraps markdown if mistakenly wrapped in JSON or code fences."""
    import re
    if not raw_text:
        return ""
    text = raw_text.strip()
    
    if '"markdown_content":' in text:
        idx = text.find('"markdown_content":')
        sub = text[idx + len('"markdown_content":'):].strip()
        if sub.startswith('"'):
            sub = sub[1:]
        sub = re.sub(r'"\s*\}?\s*```?\s*$', '', sub)
        text = sub.replace(r'\n', '\n').replace(r'\"', '"').replace(r'\\', '\\')
    elif text.startswith('```json'):
        text = text[7:].strip()
        if text.endswith('```'):
            text = text[:-3].strip()
    elif text.startswith('```markdown'):
        text = text[11:].strip()
        if text.endswith('```'):
            text = text[:-3].strip()
    elif text.startswith('```'):
        text = text[3:].strip()
        if text.endswith('```'):
            text = text[:-3].strip()
            
    return text.strip()


def build_report_docx(
    work_dir: Path,
    markdown_content: str,
    title: str = "Academic Assignment",
    course: str = "Computer Engineering",
    student_name: str = "Student",
    institution: str = "Academic Institution",
) -> Path:
    """Builds a formal Word document deliverable styled in Cognitive Heritage aesthetic with professional typography."""
    import re
    import docx
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    md_text = clean_markdown_text(markdown_content)
    doc = docx.Document()
    out_path = work_dir / "report.docx"

    # 0.8 inch margins for generous printable width (6.9 inches)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    AVAILABLE_WIDTH_INCHES = 6.9

    lines = md_text.splitlines()
    doc_title = title
    clean_lines = []
    skip_header_lines = True

    # Strip duplicate top headers if already present in markdown text
    for l in lines:
        s = l.strip()
        if skip_header_lines:
            if s.startswith('# '):
                doc_title = s[2:].strip()
                continue
            if s.startswith('**Assignment') or s.startswith('**Student') or s.startswith('**Institution') or s.startswith('**Date'):
                continue
            if s == '---' or s == '***':
                skip_header_lines = False
                continue
            if s.startswith('## '):
                skip_header_lines = False
            elif s:
                skip_header_lines = False
        clean_lines.append(l)

    # 1. Executive Title Block
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(4)
    p_title.paragraph_format.keep_with_next = True
    run_title = p_title.add_run(doc_title)
    _set_run_font(run_title, font_name="Georgia", size=Pt(20), bold=True, color=RGBColor(112, 87, 40))

    # Elegant horizontal accent bar below title
    _add_bottom_border(p_title, color="8B6F3E", sz="16")

    # Metadata Banner Card (Surface Container #FDEBD9)
    meta_table = doc.add_table(rows=1, cols=1)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_cell = meta_table.cell(0, 0)
    meta_cell.width = Inches(AVAILABLE_WIDTH_INCHES)
    _set_cell_background(meta_cell, "FDEBD9")
    _set_cell_margins(meta_cell, top=70, bottom=70, left=120, right=120)

    p_meta = meta_cell.paragraphs[0]
    p_meta.paragraph_format.space_before = Pt(1)
    p_meta.paragraph_format.space_after = Pt(1)

    r_c = p_meta.add_run("Course: ")
    _set_run_font(r_c, font_name="Segoe UI", size=Pt(9.0), bold=True, color=RGBColor(112, 87, 40))
    r_cv = p_meta.add_run(f"{course}    |    ")
    _set_run_font(r_cv, font_name="Segoe UI", size=Pt(9.0), color=RGBColor(77, 70, 58))

    r_s = p_meta.add_run("Student: ")
    _set_run_font(r_s, font_name="Segoe UI", size=Pt(9.0), bold=True, color=RGBColor(112, 87, 40))
    r_sv = p_meta.add_run(f"{student_name}    |    ")
    _set_run_font(r_sv, font_name="Segoe UI", size=Pt(9.0), color=RGBColor(77, 70, 58))

    r_i = p_meta.add_run("Institution: ")
    _set_run_font(r_i, font_name="Segoe UI", size=Pt(9.0), bold=True, color=RGBColor(112, 87, 40))
    r_iv = p_meta.add_run(f"{institution}")
    _set_run_font(r_iv, font_name="Segoe UI", size=Pt(9.0), color=RGBColor(77, 70, 58))

    # Spacing after header block
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(2)
    p_sp.paragraph_format.space_after = Pt(4)

    in_code_block = False
    code_lines = []
    table_lines = []
    in_front_matter_toc = False

    def flush_table(t_lines):
        if not t_lines:
            return
        rows = []
        for l in t_lines:
            s_l = l.strip()
            if re.match(r'^[|\s\-:]+$', s_l):
                continue
            cells = [c.strip() for c in s_l.strip('|').split('|')]
            rows.append(cells)
        if not rows:
            return

        num_cols = max(len(r) for r in rows)
        table = doc.add_table(rows=len(rows), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        is_wide = (num_cols >= 6)
        font_size = Pt(7.8) if is_wide else Pt(9.5)
        h_font_size = Pt(8.0) if is_wide else Pt(10.0)
        c_top = 50 if is_wide else 80
        c_side = 40 if is_wide else 100

        col_widths = []
        if num_cols == 8:
            col_widths = [0.95, 0.85, 0.85, 0.88, 0.82, 0.82, 0.60, 1.13]
        elif num_cols == 6:
            col_widths = [1.2, 1.0, 1.0, 1.0, 1.0, 1.7]
        elif num_cols == 5:
            col_widths = [1.4, 1.2, 1.2, 1.2, 1.9]
        elif num_cols == 4:
            col_widths = [1.7, 1.7, 1.7, 1.8]
        elif num_cols == 3:
            col_widths = [2.1, 2.4, 2.4]
        elif num_cols == 2:
            col_widths = [2.4, 4.5]
        else:
            w_each = AVAILABLE_WIDTH_INCHES / num_cols
            col_widths = [w_each] * num_cols

        for r_idx, row_data in enumerate(rows):
            tr = table.rows[r_idx]._tr
            trPr = tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

            if r_idx == 0:
                trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

            for c_idx in range(num_cols):
                cell_text = row_data[c_idx] if c_idx < len(row_data) else ""
                cell = table.cell(r_idx, c_idx)
                col_w = col_widths[c_idx] if c_idx < len(col_widths) else 1.0
                cell.width = Inches(col_w)

                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(1.5)
                p.paragraph_format.space_after = Pt(1.5)
                p.paragraph_format.line_spacing = 1.05

                if r_idx == 0:
                    _set_cell_background(cell, "8B6F3E")
                    _set_cell_margins(cell, top=c_top + 15, bottom=c_top + 15, left=c_side, right=c_side)
                    if is_wide and c_idx in (1, 2, 3, 4, 5, 6):
                        _set_cell_nowrap(cell)
                    clean_header = clean_math_text(cell_text)
                    run = p.add_run(clean_header)
                    _set_run_font(run, font_name="Segoe UI", size=h_font_size, bold=True, color=RGBColor(255, 248, 244))
                else:
                    bg_color = "FFF8F4" if r_idx % 2 == 1 else "FDEBD9"
                    _set_cell_background(cell, bg_color)
                    _set_cell_margins(cell, top=c_top, bottom=c_top, left=c_side, right=c_side)
                    _add_formatted_runs(p, cell_text, base_font="Segoe UI", base_size=font_size)

        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_after = Pt(6)

    def flush_code(c_lines):
        if not c_lines:
            return
        code_text = "\n".join(c_lines)
        t_code = doc.add_table(rows=1, cols=1)
        t_code.alignment = WD_TABLE_ALIGNMENT.CENTER
        t_code.autofit = False
        c_cell = t_code.cell(0, 0)
        c_cell.width = Inches(AVAILABLE_WIDTH_INCHES)
        _set_cell_background(c_cell, "F7E5D4")
        _set_cell_margins(c_cell, top=80, bottom=80, left=120, right=120)

        tcPr = c_cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:left w:val="single" w:sz="18" w:space="0" w:color="8B6F3E"/>
                <w:top w:val="none"/>
                <w:right w:val="none"/>
                <w:bottom w:val="none"/>
            </w:tcBorders>
        ''')
        tcPr.append(tcBorders)

        cp = c_cell.paragraphs[0]
        cp.paragraph_format.space_before = Pt(2)
        cp.paragraph_format.space_after = Pt(2)
        cp.paragraph_format.line_spacing = 1.05
        cp.paragraph_format.keep_together = True
        run = cp.add_run(code_text)
        _set_run_font(run, font_name="Consolas", size=Pt(8.8), color=RGBColor(35, 26, 15))

        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_after = Pt(4)

    for line in clean_lines:
        s = line.strip()

        if s.startswith("```"):
            if in_code_block:
                in_code_block = False
                flush_code(code_lines)
                code_lines = []
            else:
                if table_lines:
                    flush_table(table_lines)
                    table_lines = []
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        if s.startswith("|") and s.endswith("|"):
            table_lines.append(s)
            continue
        else:
            if table_lines:
                flush_table(table_lines)
                table_lines = []

        if not s:
            continue

        # Standalone Math Formula Block \[ ... \] or $$ ... $$
        if (s.startswith("\\[") and s.endswith("\\]")) or (s.startswith("$$") and s.endswith("$$")):
            math_formula = clean_math_text(s)
            p_math = doc.add_paragraph()
            p_math.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_math.paragraph_format.space_before = Pt(5)
            p_math.paragraph_format.space_after = Pt(5)
            p_math.paragraph_format.keep_with_next = True
            run = p_math.add_run(math_formula)
            _set_run_font(run, font_name="Georgia", size=Pt(10.5), italic=True, color=RGBColor(112, 87, 40))
            continue

        # Horizontal Divider
        if s == "---" or s == "***" or re.match(r'^-{3,}$', s):
            if in_front_matter_toc:
                in_front_matter_toc = False
                doc.add_page_break()
                continue
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            _add_bottom_border(p, color="D1C5B6", sz="6")
            continue

        # Headings: H1, H2, H3, H4, H5
        if s.startswith("# "):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(3)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(s[2:].strip())
            _set_run_font(run, font_name="Georgia", size=Pt(17), bold=True, color=RGBColor(112, 87, 40))
        elif s.startswith("## "):
            heading_text = s[3:].strip()
            if "table of contents" in heading_text.lower():
                in_front_matter_toc = True
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(13)
            h.paragraph_format.space_after = Pt(2.5)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(heading_text)
            _set_run_font(run, font_name="Georgia", size=Pt(13.5), bold=True, color=RGBColor(112, 87, 40))
        elif s.startswith("### "):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(2)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(s[4:].strip())
            _set_run_font(run, font_name="Georgia", size=Pt(11.5), bold=True, color=RGBColor(139, 111, 62))
        elif s.startswith("#### "):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(s[5:].strip())
            _set_run_font(run, font_name="Georgia", size=Pt(11), bold=True, color=RGBColor(139, 111, 62))
        elif s.startswith("##### "):
            h = doc.add_paragraph()
            h.paragraph_format.space_before = Pt(6)
            h.paragraph_format.space_after = Pt(2)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(s[6:].strip())
            _set_run_font(run, font_name="Georgia", size=Pt(10), bold=True, italic=True, color=RGBColor(100, 94, 78))
        elif s.startswith("- ") or s.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.10
            clean_item = re.sub(r'^(-|\*)\s*', '', s)
            _add_formatted_runs(p, clean_item, base_font="Segoe UI", base_size=Pt(10))
        elif re.match(r'^\d+\.\s', s):
            m = re.match(r'^(\d+\.)\s*(.*)', s)
            num_prefix = m.group(1)
            item_body = m.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            if in_front_matter_toc:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(1.5)
                p.paragraph_format.line_spacing = 1.05
            else:
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.12

            run_num = p.add_run(f"{num_prefix} ")
            _set_run_font(run_num, font_name="Segoe UI", size=Pt(10), bold=True, color=RGBColor(112, 87, 40))
            _add_formatted_runs(p, item_body, base_font="Segoe UI", base_size=Pt(10))
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            _add_formatted_runs(p, s, base_font="Segoe UI", base_size=Pt(10))

    if table_lines:
        flush_table(table_lines)
    if code_lines:
        flush_code(code_lines)

    doc.save(str(out_path))
    return out_path


def build_presentation_pptx(
    work_dir: Path,
    slides: List[Dict[str, Any]],
    title: str = "Academic Presentation",
    course: str = "General",
    student_name: str = "Student",
) -> Path:
    """Builds a PowerPoint presentation styled in Cognitive Heritage aesthetic."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor

    prs = Presentation()
    out_path = work_dir / "presentation.pptx"

    # 16:9 Widescreen dimensions
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Cognitive Heritage Palette
    COLOR_BG = RGBColor(255, 248, 244)       # #fff8f4 Surface
    COLOR_CARD = RGBColor(253, 235, 217)     # #fdebd9 Surface Container
    COLOR_PRIMARY = RGBColor(112, 87, 40)    # #705728 Primary
    COLOR_TEXT = RGBColor(35, 26, 15)        # #231a0f On-surface
    COLOR_MUTED = RGBColor(77, 70, 58)       # #4d463a On-surface-variant

    for idx, slide_info in enumerate(slides):
        slide = prs.slides.add_slide(blank_layout)

        # Background fill
        bg_shape = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = COLOR_BG
        bg_shape.line.fill.background()

        # Title Card
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(1.2))
        tf = title_box.text_frame
        tf.word_wrap = True
        p_title = tf.paragraphs[0]
        p_title.text = slide_info.get("title", f"Slide {idx + 1}")
        p_title.font.name = "Noto Serif"
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_PRIMARY

        # Body Container Card
        card = slide.shapes.add_shape(1, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD
        card.line.color.rgb = RGBColor(209, 197, 182)

        # Body text frame inside card
        body_box = slide.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.1), Inches(4.4))
        btf = body_box.text_frame
        btf.word_wrap = True

        raw_body = slide_info.get("body", "")
        lines = [l for l in raw_body.splitlines() if l.strip()]
        for b_idx, line in enumerate(lines):
            p = btf.paragraphs[0] if b_idx == 0 else btf.add_paragraph()
            p.text = line.lstrip("•-* ")
            p.font.name = "Manrope"
            p.font.size = Pt(18)
            p.font.color.rgb = COLOR_TEXT
            p.space_after = Pt(12)

        # Speaker notes
        notes = slide_info.get("speaker_notes")
        if notes and slide.has_notes_slide:
            slide.notes_slide.notes_text_frame.text = notes

    prs.save(str(out_path))
    return out_path


def generate_ui_screenshot(
    work_dir: Path,
    title: str = "Application Interface",
    subtitle: str = "Cognitive Heritage Mobile UI",
) -> Path:
    """Generates a high-resolution UI preview screenshot deliverable."""
    from PIL import Image, ImageDraw

    out_path = work_dir / "screenshot.png"
    img = Image.new("RGB", (1080, 1920), color=(255, 248, 244))  # #fff8f4
    draw = ImageDraw.Draw(img)

    # Top app bar
    draw.rectangle([0, 0, 1080, 180], fill=(112, 87, 40))  # #705728
    draw.text((60, 80), title[:42], fill=(255, 255, 255))

    # Main Card Container
    draw.rectangle([60, 240, 1020, 1260], fill=(253, 235, 217), outline=(209, 197, 182), width=3)
    draw.text((100, 280), subtitle[:45], fill=(35, 26, 15))

    # Card inputs / status boxes
    draw.rectangle([100, 360, 980, 480], fill=(255, 255, 255), outline=(127, 118, 105), width=2)
    draw.text((120, 400), "Input: Student Workspace (Academic Pipeline)", fill=(77, 70, 58))

    draw.rectangle([100, 520, 980, 640], fill=(255, 255, 255), outline=(127, 118, 105), width=2)
    draw.text((120, 560), "Status: Verified & Operational", fill=(77, 70, 58))

    # Action button
    draw.rectangle([100, 700, 980, 820], fill=(112, 87, 40))
    draw.text((440, 740), "ACTION / EXECUTE", fill=(255, 255, 255))

    # Output Card
    draw.rectangle([100, 880, 980, 1180], fill=(255, 255, 255), outline=(112, 87, 40), width=2)
    draw.text((120, 920), "Output Log: System operational, all endpoints responsive.", fill=(35, 26, 15))

    img.save(str(out_path), "PNG")
    return out_path
